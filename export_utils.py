import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def generate_styled_distance_matrix_excel(distance_matrix):
    """
    รับข้อมูล distance_matrix จริงจากโปรเจกต์ มาสร้างไฟล์ Excel ตกแต่งสวยงาม
    และส่งคืนเป็นข้อมูลไบต์ (BytesIO) สำหรับใช้ในปุ่มดาวน์โหลดของ Streamlit
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Distance Matrix"

    num_places = len(distance_matrix)
    
    # 1. เพิ่มหัวตาราง
    headers = ["สถานที่"] + [str(i) for i in range(1, num_places + 1)]
    ws.append(headers)

    # 2. เพิ่มข้อมูลลงตาราง (ปัดเศษ 2 ตำแหน่งให้สวยงาม)
    for i, row_data in enumerate(distance_matrix):
        rounded_row = [round(val, 2) for val in row_data]
        row = [str(i + 1)] + rounded_row
        ws.append(row)

    # 3. กำหนดสไตล์ให้เหมือนรูปภาพ
    blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    grey_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal="center", vertical="center")

    # 4. ไล่ปรับสไตล์ทุกเซลล์
    for row in ws.iter_rows(min_row=1, max_row=num_places + 1, min_col=1, max_col=num_places + 1):
        for cell in row:
            cell.border = thin_border
            cell.alignment = center_align
            
            # หัวคอลัมน์และหัวแถวเป็นสีฟ้าอ่อน
            if cell.row == 1 or cell.column == 1:
                cell.fill = blue_fill
                cell.font = Font(bold=True)
                
            # แนวทแยงมุม (ระยะทางจากตัวเองไปตัวเอง = 0) เป็นสีเทา
            if cell.row > 1 and cell.column > 1 and cell.row == cell.column:
                cell.fill = grey_fill

    # 5. ปรับความกว้างของคอลัมน์
    ws.column_dimensions['A'].width = 12
    for col in range(2, num_places + 2):
        # รองรับคอลัมน์มากกว่า Z (A..Z, AA..AZ)
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = 8

    # 6. บันทึกลงหน่วยความจำ (ไม่ต้องสร้างไฟล์ขยะในเครื่อง)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output
    