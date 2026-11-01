import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

UNREACHABLE_HOURS = 99.0    # optimiser.py แทนค่า inf ด้วย 100 ชม. / 5000 กม. (ถือว่าเข้าไม่ถึง)
UNREACHABLE_KM = 4999.0

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
DIAG_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _add_sheet(ws, matrix, labels, number_format, unreachable_at):
    n = len(matrix)
    ws.cell(row=1, column=1, value="จาก \\ ไป")
    for j in range(n):
        ws.cell(row=1, column=j + 2, value=labels[j])
    for i in range(n):
        ws.cell(row=i + 2, column=1, value=labels[i])
        for j in range(n):
            val = matrix[i][j]
            cell = ws.cell(row=i + 2, column=j + 2)
            if i == j:
                cell.value = 0
                cell.fill = DIAG_FILL
            elif val is None or val == float("inf") or val >= unreachable_at:
                cell.value = "เข้าไม่ถึง"
            else:
                cell.value = val
            cell.number_format = number_format
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER

    # หัวตาราง
    for c in range(1, n + 2):
        for r in (1,) if c > 1 else range(1, n + 2):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
    ws.row_dimensions[1].height = 60
    ws.column_dimensions["A"].width = 28
    for c in range(2, n + 2):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.freeze_panes = "B2"

    # ไล่สี เขียว(ใกล้) -> เหลือง -> แดง(ไกล) เฉพาะช่องตัวเลข
    rng = f"B2:{get_column_letter(n + 1)}{n + 1}"
    ws.conditional_formatting.add(
        rng,
        ColorScaleRule(start_type="min", start_color="C6EFCE",
                       mid_type="percentile", mid_value=50, mid_color="FFEB9C",
                       end_type="max", end_color="F8CBAD"),
    )


def generate_styled_matrices_excel(dist_matrix, time_matrix, names=None):
    """สร้างไฟล์ Excel 3 ชีต: ระยะทาง (กม.), เวลาเดินทาง (นาที), เวลาเดินทาง (ชั่วโมง)
    dist_matrix: กม., time_matrix: ชั่วโมง (รูปแบบเดียวกับที่ get_travel_matrices คืนมา)
    names: ชื่อสถานที่ตามลำดับ index (ไม่ใส่ก็ใช้เลข 1..n)
    คืนค่าเป็น bytes ใช้กับ st.download_button ได้เลย"""
    n = len(dist_matrix)
    if names:
        labels = [f"{i + 1}. {names[i]}" for i in range(n)]
    else:
        labels = [str(i + 1) for i in range(n)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Distance (km)"
    _add_sheet(ws, dist_matrix, labels, "0.00", UNREACHABLE_KM)

    minutes = [[(x * 60 if x is not None and x != float("inf") else x) for x in row] for row in time_matrix]
    _add_sheet(wb.create_sheet("Time (min)"), minutes, labels, "0.0", UNREACHABLE_HOURS * 60)
    _add_sheet(wb.create_sheet("Time (hr)"), time_matrix, labels, "0.00", UNREACHABLE_HOURS)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()